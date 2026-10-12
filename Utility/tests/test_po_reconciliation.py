from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "Utility" / "src"))

from config_loader import load_json
from database_service import ReconciliationDatabase
from pipeline import run_inventory_scan
from reconciliation_engine import _po_facility_matches, build_and_store_lifecycle, build_lifecycle_rows
from report_generator import build_report


def insert_record(database, table, **values):
    """Small real-SQLite fixtures, with required lineage fields filled in."""
    for column in database.connect().execute(f"PRAGMA table_info({table})"):
        _, name, kind, required, default, primary = column
        if required and default is None and not primary and name not in values:
            values[name] = 1 if kind == "INTEGER" else "fixture"
    columns = ", ".join(values)
    placeholders = ", ".join("?" for _ in values)
    database.connect().execute(f"INSERT INTO {table} ({columns}) VALUES ({placeholders})", list(values.values()))
    database.connect().commit()


class FacilityTests(unittest.TestCase):
    def test_normalized_same_hospital(self):
        for location in ["AH Choppadandi", "A.H. Choppadandi", "Area Hospital Choppadandi",
                         "AH Choppadandi Dept. of Neonatalogy", "AH Choppadandi Department of Cardiology"]:
            with self.subTest(location=location):
                self.assertTrue(_po_facility_matches("AH, Choppadandi", "Karimnagar", location, " Karimnagar District "))

    def test_other_hospitals_and_conflicting_districts_rejected(self):
        for location, district in [("AH Nakrekal", "Nalgonda"), ("AH Nakrekal", "Karimnagar"),
                                   ("DH Choppadandi", "Karimnagar"), ("AH Choppadandi", "Nalgonda")]:
            with self.subTest(location=location, district=district):
                self.assertFalse(_po_facility_matches("AH, Choppadandi", "Karimnagar", location, district))

    def test_unknown_is_not_a_wildcard(self):
        for missing in [None, "", "nan", float("nan"), "-", "<NA>"]:
            self.assertFalse(_po_facility_matches("AH Choppadandi", "Karimnagar", missing, "Karimnagar"))
            self.assertFalse(_po_facility_matches(missing, "Karimnagar", missing, "Karimnagar"))
            self.assertFalse(_po_facility_matches("AH Choppadandi", missing, "AH Choppadandi", missing))


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db = ReconciliationDatabase(Path(self.temp.name) / "test.db")
        self.db.initialize()
        insert_record(self.db, "source_files", file_id=1, source_type="Indent", file_hash="indent-hash")
        insert_record(self.db, "source_sheets", sheet_id=1, file_id=1)
        insert_record(self.db, "indent_headers", indent_header_id=1, institution_name="AH, Choppadandi", district="Karimnagar")
        self.rules = {"item_name_threshold": 95, "item_probable_threshold": 70, "approved_aliases": {"Syrange Pump": "Syringe Pump"}}

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def indent(self, item="Infusion Pump", qty="2", line_id=1):
        insert_record(self.db, "indent_lines", indent_line_id=line_id, indent_header_id=1,
                      source_item_name=item, normalized_item_name=item, required_qty=qty)

    def po(self, item="Infusion Pump", qty="2", po_id=1, number="215A", location="AH Choppadandi", district="Karimnagar"):
        insert_record(self.db, "po_lines", po_line_id=po_id, po_no=number, source_item_name=item,
                      ordered_qty=qty, location=location, district=district)

    def test_valid_po_restored_foreign_po_excluded(self):
        self.indent()
        self.po(po_id=1, number="345A", location="AH Nakrekal", district="Nalgonda")
        self.po(po_id=2)
        rows, allocations = build_lifecycle_rows(self.db, self.rules)
        self.assertEqual(json.loads(rows[0].matched_po_ids), [2])
        self.assertEqual([a.po_no for a in allocations], ["215A"])
        self.assertEqual(rows[0].confirmed_ordered_qty, "2")

    def test_approved_alias_and_punctuation(self):
        for indent, po in [("Syrange Pump", "Syringe Pump"), ("Crashcart", "Crash cart"),
                           ("Operation Table-Hydraulic Major", "Operation Table Hydraulic Major")]:
            with self.subTest(item=indent):
                self.db.connect().execute("DELETE FROM indent_lines")
                self.db.connect().execute("DELETE FROM po_lines")
                self.indent(item=indent)
                self.po(item=po)
                self.assertEqual(len(build_lifecycle_rows(self.db, self.rules)[1]), 1)

    def test_weak_or_conflicting_items_are_not_confirmed(self):
        for indent, po in [("CTG Machines", "12 Channel ECG Machine"),
                           ("100 MA X-ray Machine", "200 MA X-ray Machine"),
                           ("MRI 1.5 Tesla", "MRI 15 Tesla"),
                           ("Monitors (3 Para)", "Multipara Monitor (7 Para)"),
                           ("Dental Chair", "Dental X-Ray-Intra Oral With RVG"),
                           ("Operation table Hydraulic Minor", "Operation Table Hydraulic Major")]:
            with self.subTest(item=indent):
                self.db.connect().execute("DELETE FROM indent_lines")
                self.db.connect().execute("DELETE FROM po_lines")
                self.indent(item=indent)
                self.po(item=po)
                rows, allocations = build_lifecycle_rows(self.db, self.rules)
                self.assertEqual(allocations, [])
                self.assertEqual(rows[0].matched_po_ids, "[]")
                self.assertEqual(rows[0].confirmed_ordered_qty, "0")
                self.assertTrue(rows[0].raw_payload["po_review_candidates"])

    def test_semantic_keyword_guard_rejects_unrelated_items(self):
        """With the real keyword/attribute guard config applied, items that
        merely share a generic noun (e.g. "machine") but are semantically
        unrelated must not be surfaced as probable candidates at all."""
        matching_rules_path = ROOT / "Utility" / "config" / "MatchingRules.json"
        real_rules = load_json(matching_rules_path)
        rules = dict(self.rules)
        rules["generic_stopwords"] = real_rules["generic_stopwords"]
        rules["attribute_keyword_map"] = real_rules["attribute_keyword_map"]
        for indent, po in [("Weighing machine adult", "12 Channel ECG Machine"),
                           ("12 Channel ECG Machine", "Weighing machine adult")]:
            with self.subTest(indent=indent, po=po):
                self.db.connect().execute("DELETE FROM indent_lines")
                self.db.connect().execute("DELETE FROM po_lines")
                self.indent(item=indent)
                self.po(item=po)
                rows, allocations = build_lifecycle_rows(self.db, rules)
                self.assertEqual(allocations, [])
                self.assertEqual(rows[0].matched_po_ids, "[]")
                self.assertEqual(rows[0].raw_payload["po_review_candidates"], [])

    def test_approved_substitution_bypasses_numeric_guard(self):
        """An alias that maps two differently-numbered labels to the same
        canonical equipment (e.g. an approved 7-Para monitor instead of a
        5-Para one) must confirm, unlike an incidental digit collision."""
        rules = dict(self.rules)
        rules["approved_aliases"] = {
            **self.rules["approved_aliases"],
            "Multipara Monitor (7 Para)": "Monitors ( 5 Para)",
        }
        self.indent(item="Monitors ( 5 Para)", qty="9")
        self.po(item="Multipara Monitor (7 Para)", qty="9")
        rows, allocations = build_lifecycle_rows(self.db, rules)
        self.assertEqual(len(allocations), 1)
        self.assertEqual(rows[0].confirmed_ordered_qty, "9")
        self.assertEqual(rows[0].po_pending_qty, "0")

    def test_package_component_attached_without_affecting_quantity(self):
        """Accessory PO lines bundled with a matched anchor item (e.g. CR
        imaging plates with a CR machine) are linked for visibility but do not
        inflate the indent's confirmed/allocated quantity."""
        rules = dict(self.rules)
        rules["approved_aliases"] = {
            **self.rules["approved_aliases"],
            "Computed Radiography Machine : CR Imaging pates and compatible cassettes:": "CR Machine",
        }
        rules["package_components"] = {
            "216A": {
                "anchor": "Computed Radiography Machine : CR Imaging pates and compatible cassettes:",
                "components": ["35 cm x 43 cm plate", "24 cm x 30 cm plate"],
            }
        }
        self.indent(item="CR Machine", qty="1")
        self.po(item="Computed Radiography Machine : CR Imaging pates and compatible cassettes:", po_id=1, number="216A", qty="1")
        self.po(item="35 cm x 43 cm plate", po_id=2, number="216A", qty="1")
        self.po(item="24 cm x 30 cm plate", po_id=3, number="216A", qty="1")
        rows, allocations = build_lifecycle_rows(self.db, rules)
        self.assertEqual(sorted(json.loads(rows[0].matched_po_ids)), [1, 2, 3])
        self.assertEqual(rows[0].confirmed_ordered_qty, "1")
        self.assertEqual(rows[0].po_pending_qty, "0")
        component_statuses = sorted(a.allocation_status for a in allocations)
        self.assertEqual(component_statuses, ["Allocated", "Package Component", "Package Component"])

    def test_unmatched_po_line_without_indent_counterpart_is_recorded_as_exception(self):
        self.indent(item="Office Table", qty="3", line_id=1)
        self.po(item="Office Table", po_id=1, number="218A", qty="3")
        self.po(item="Patient Weighing Scale", po_id=2, number="218A", qty="7")
        build_and_store_lifecycle(self.db, self.rules)
        exceptions = self.db.fetch_dataframe("SELECT exception_code, row_id FROM exceptions")
        orphan = exceptions[exceptions["exception_code"] == "PO_LINE_NOT_FOUND_IN_INDENT"]
        self.assertEqual(orphan["row_id"].tolist(), [2])

    def test_po_header_amount_mismatch_is_flagged(self):
        self.indent()
        insert_record(self.db, "po_lines", po_line_id=1, po_no="223A", source_item_name="Stretcher on Trolley",
                      ordered_qty="8", location="AH Choppadandi", district="Karimnagar",
                      gross_amount="29958", total_amount="217072")
        insert_record(self.db, "po_lines", po_line_id=2, po_no="223A", source_item_name="Saline Stand",
                      ordered_qty="20", location="AH Choppadandi", district="Karimnagar",
                      gross_amount="29958", total_amount="56480")
        build_and_store_lifecycle(self.db, self.rules)
        exceptions = self.db.fetch_dataframe("SELECT exception_code, summary FROM exceptions WHERE exception_code='PO_HEADER_AMOUNT_MISMATCH'")
        self.assertEqual(len(exceptions), 1)
        self.assertIn("223A", exceptions.iloc[0]["summary"])

    def test_rerun_does_not_duplicate_data_quality_exceptions(self):
        self.indent(item="Office Table", qty="3", line_id=1)
        self.po(item="Office Table", po_id=1, number="218A", qty="3")
        self.po(item="Patient Weighing Scale", po_id=2, number="218A", qty="7")
        build_and_store_lifecycle(self.db, self.rules)
        build_and_store_lifecycle(self.db, self.rules)
        exceptions = self.db.fetch_dataframe("SELECT exception_code FROM exceptions WHERE exception_code='PO_LINE_NOT_FOUND_IN_INDENT'")
        self.assertEqual(len(exceptions), 1)

    def test_excess_po_quantity_over_required_is_flagged(self):
        self.indent(item="X-ray Film", qty="5", line_id=1)
        self.po(item="X-ray Film", po_id=1, number="219A", qty="10")
        build_and_store_lifecycle(self.db, self.rules)
        exceptions = self.db.fetch_dataframe(
            "SELECT exception_code, details_json FROM exceptions WHERE exception_code='PO_QUANTITY_EXCEEDS_REQUIRED'"
        )
        self.assertEqual(len(exceptions), 1)
        details = json.loads(exceptions.iloc[0]["details_json"])
        self.assertEqual(details["required_qty"], 5)
        self.assertEqual(details["candidate_qty"], 10)
        self.assertEqual(details["excess_qty"], 5)

    def test_rerun_does_not_duplicate_excess_quantity_exceptions(self):
        self.indent(item="X-ray Film", qty="5", line_id=1)
        self.po(item="X-ray Film", po_id=1, number="219A", qty="10")
        build_and_store_lifecycle(self.db, self.rules)
        build_and_store_lifecycle(self.db, self.rules)
        exceptions = self.db.fetch_dataframe(
            "SELECT exception_code FROM exceptions WHERE exception_code='PO_QUANTITY_EXCEEDS_REQUIRED'"
        )
        self.assertEqual(len(exceptions), 1)

    def test_duplicate_po_lines_at_same_facility_are_flagged(self):
        self.indent(item="Stretcher", qty="5", line_id=1)
        self.po(item="Stretcher", po_id=1, number="400A", qty="5",
                location="AH Choppadandi", district="Karimnagar")
        self.po(item="Stretcher", po_id=2, number="400A", qty="5",
                location="AH Choppadandi", district="Karimnagar")
        self.po(item="Stretcher", po_id=3, number="400A", qty="5",
                location="AH Nakrekal", district="Nalgonda")
        build_and_store_lifecycle(self.db, self.rules)
        exceptions = self.db.fetch_dataframe(
            "SELECT exception_code, details_json FROM exceptions WHERE exception_code='DUPLICATE_PO_LINE'"
        )
        self.assertEqual(len(exceptions), 1)
        details = json.loads(exceptions.iloc[0]["details_json"])
        self.assertEqual(details["duplicate_count"], 2)
        self.assertEqual(sorted(details["po_line_ids"]), [1, 2])

    def test_quantity_is_not_reused_across_indent_lines(self):
        self.indent(qty="2", line_id=1)
        self.indent(qty="2", line_id=2)
        self.po(qty="3")
        rows, allocations = build_lifecycle_rows(self.db, self.rules)
        self.assertEqual(sum(float(a.allocated_qty) for a in allocations), 3)
        self.assertEqual([r.po_pending_qty for r in rows], ["0", "1"])

    def test_probable_po_line_is_not_offered_to_more_than_one_indent_row(self):
        """A single PO line that fuzzy-matches multiple unrelated indent rows
        must only appear as a probable candidate for one of them - it has
        only one physical line row to offer."""
        self.indent(item="Patient Monitor", qty="2", line_id=1)
        self.indent(item="Multipara Monitor", qty="2", line_id=2)
        self.po(item="Patient Monitor Multipara", po_id=1, number="300A", qty="5")
        rows, allocations = build_lifecycle_rows(self.db, self.rules)
        rows = sorted(rows, key=lambda record: record.indent_line_id)
        review_po_ids_by_line = {
            record.indent_line_id: {c["po_line_id"] for c in record.raw_payload["po_review_candidates"]}
            for record in rows
        }
        claimed_by = [line_id for line_id, po_ids in review_po_ids_by_line.items() if 1 in po_ids]
        self.assertEqual(len(claimed_by), 1)

    def test_po_number_reuse_capped_to_its_line_count(self):
        """A PO number with exactly 2 line rows must never show up as a
        candidate for more than 2 distinct indent lines, even when 3+ indent
        lines are plausible fuzzy matches for its items."""
        self.indent(item="Multipara Monitor X", qty="1", line_id=10)
        self.indent(item="Multipara Monitor Y", qty="1", line_id=11)
        self.indent(item="Multipara Monitor Z", qty="1", line_id=12)
        self.po(item="Multipara Monitor A", po_id=1, number="301A", qty="2")
        self.po(item="Multipara Monitor B", po_id=2, number="301A", qty="2")
        rows, allocations = build_lifecycle_rows(self.db, self.rules)
        for po_line_id in (1, 2):
            claiming_lines = [
                record.indent_line_id
                for record in rows
                if any(c["po_line_id"] == po_line_id for c in record.raw_payload["po_review_candidates"])
            ]
            self.assertLessEqual(len(claiming_lines), 1)
        lines_showing_po_no = {
            record.indent_line_id
            for record in rows
            if any(c["po_no"] == "301A" for c in record.raw_payload["po_review_candidates"])
        }
        self.assertLessEqual(len(lines_showing_po_no), 2)

    def test_vendor_split_allocation_unaffected_by_reuse_cap(self):
        """One indent line's quantity split across multiple different PO
        numbers/vendors must still fully allocate - the reuse cap only
        prevents the SAME po_line_id being claimed twice, never limits how
        many different PO lines one indent row can draw from."""
        self.indent(qty="5")
        self.po(qty="2", po_id=1, number="215A")
        self.po(qty="3", po_id=2, number="216A")
        rows, allocations = build_lifecycle_rows(self.db, self.rules)
        self.assertEqual(json.loads(rows[0].matched_po_ids), [1, 2])
        self.assertEqual(sum(float(a.allocated_qty) for a in allocations), 5)
        self.assertEqual(rows[0].po_pending_qty, "0")

    def test_exact_match_has_priority_over_earlier_fuzzy(self):
        self.indent(item="Infusion Pumps", line_id=1)
        self.indent(item="Infusion Pump", line_id=2)
        self.po()
        _, allocations = build_lifecycle_rows(self.db, self.rules)
        self.assertEqual([a.indent_line_id for a in allocations], [2])

    def test_split_po_quantities(self):
        self.indent(qty="5")
        self.po(qty="2", po_id=1, number="215A")
        self.po(qty="3", po_id=2, number="216A")
        rows, allocations = build_lifecycle_rows(self.db, self.rules)
        self.assertEqual(json.loads(rows[0].matched_po_ids), [1, 2])
        self.assertEqual(sum(float(a.allocated_qty) for a in allocations), 5)

    def test_empty_po_table_retains_schema(self):
        self.indent()
        rows, allocations = build_lifecycle_rows(self.db, self.rules)
        self.assertEqual(len(rows), 1)
        self.assertEqual(allocations, [])

    def test_confidence_score_reflects_match_quality(self):
        """A clean, fully-allocated confirmed match should score markedly
        higher than a weak/ambiguous probable-only match."""
        self.indent(item="Infusion Pump", qty="2", line_id=1)
        self.po(item="Infusion Pump", po_id=1, number="215A", qty="2")
        self.indent(item="100 MA X-ray Machine", qty="1", line_id=2)
        self.po(item="200 MA X-ray Machine", po_id=2, number="216A", qty="1")
        rows, _ = build_lifecycle_rows(self.db, self.rules)
        rows_by_line = {record.indent_line_id: record for record in rows}
        confirmed_score = rows_by_line[1].confidence_score
        probable_score = rows_by_line[2].confidence_score
        self.assertGreaterEqual(confirmed_score, 80.0)
        self.assertLess(probable_score, confirmed_score)

    def test_unlisted_sibling_po_line_is_suggested_as_component(self):
        """An unlisted sibling PO line sharing a po_no with a confirmed anchor
        and containing a known component-hint keyword should be surfaced as a
        suggested component - never silently merged into confirmed quantity."""
        rules = dict(self.rules)
        rules["component_keyword_hints"] = ["plate"]
        self.indent(item="CR Machine", qty="1", line_id=1)
        self.po(item="CR Machine", po_id=1, number="500A", qty="1")
        self.po(item="35 cm x 43 cm plate", po_id=2, number="500A", qty="1")
        rows, allocations = build_lifecycle_rows(self.db, rules)
        self.assertEqual(json.loads(rows[0].matched_po_ids), [1])
        self.assertEqual(rows[0].confirmed_ordered_qty, "1")
        suggestions = rows[0].raw_payload.get("suggested_components", [])
        self.assertEqual(len(suggestions), 1)
        self.assertEqual(suggestions[0]["po_line_id"], 2)
        self.assertEqual(suggestions[0]["status"], "Suggested Component - Confirmation Required")

    def test_suggested_component_is_not_flagged_as_orphan(self):
        """A sibling PO line surfaced as a suggested component must not also
        be reported as an unmatched/orphan PO line exception."""
        rules = dict(self.rules)
        rules["component_keyword_hints"] = ["plate"]
        self.indent(item="CR Machine", qty="1", line_id=1)
        self.po(item="CR Machine", po_id=1, number="500A", qty="1")
        self.po(item="35 cm x 43 cm plate", po_id=2, number="500A", qty="1")
        build_and_store_lifecycle(self.db, rules)
        exceptions = self.db.fetch_dataframe(
            "SELECT exception_code, row_id FROM exceptions WHERE exception_code='PO_LINE_NOT_FOUND_IN_INDENT'"
        )
        self.assertNotIn(2, exceptions["row_id"].tolist())

    def test_rebuild_replaces_stale_results_and_rolls_back_failure(self):
        self.indent()
        self.po()
        build_and_store_lifecycle(self.db, self.rules)
        build_and_store_lifecycle(self.db, self.rules)
        self.assertEqual(len(self.db.fetch_dataframe("SELECT * FROM indent_lifecycle")), 1)
        self.assertEqual(len(self.db.fetch_dataframe("SELECT * FROM po_allocations")), 1)
        rows, allocations = build_lifecycle_rows(self.db, self.rules)
        with patch.object(self.db, "record_po_allocations", side_effect=RuntimeError("test failure")):
            with self.assertRaises(RuntimeError):
                self.db.replace_reconciliation(rows, allocations)
        self.assertEqual(len(self.db.fetch_dataframe("SELECT * FROM po_allocations")), 1)
        self.db.connect().execute("UPDATE po_lines SET location='AH Nakrekal', district='Nalgonda'")
        self.db.connect().commit()
        build_and_store_lifecycle(self.db, self.rules)
        self.assertEqual(len(self.db.fetch_dataframe("SELECT * FROM po_allocations")), 0)
        self.assertEqual(self.db.fetch_dataframe("SELECT matched_po_ids FROM indent_lifecycle").iloc[0, 0], "[]")

    def test_inventory_only_import_does_not_block_extraction(self):
        self.assertIsNone(self.db.file_hash_exists("indent-hash"))
        self.indent()
        self.assertEqual(self.db.file_hash_exists("indent-hash"), 1)

    def test_locked_latest_emits_warning_but_keeps_new_report(self):
        self.indent()
        self.po()
        build_and_store_lifecycle(self.db, self.rules)
        config = {"root_path": str(ROOT / "Files"), "output_folder": self.temp.name}
        with patch("report_generator.shutil.copy2", side_effect=PermissionError):
            with self.assertWarnsRegex(RuntimeWarning, "NOT updated"):
                output = build_report(self.db, config)
        workbook = load_workbook(output)
        self.assertEqual(workbook["Indent Tracker"]["Q4"].value, "215A")
        workbook.close()


@unittest.skipUnless((ROOT / "Files" / "Purchase Orders" / "PO Register.xlsx").exists(), "Sample workbooks unavailable")
class SampleIntegrationTests(unittest.TestCase):
    def test_real_sources_and_export(self):
        with tempfile.TemporaryDirectory() as temporary:
            database = ReconciliationDatabase(Path(temporary) / "sample.db")
            database.initialize()
            config = load_json(ROOT / "Utility" / "config" / "AppConfig.json")
            config.update(root_path=str(ROOT / "Files"), output_folder=temporary, create_latest_output=False)
            try:
                first = run_inventory_scan(config, database)
                self.assertEqual(first.exceptions_created, 0)
                query = """SELECT a.po_no, a.allocated_qty, p.ordered_qty, p.po_line_id, p.location, p.district,
                           i.source_item_name, h.institution_name, h.district AS indent_district
                           FROM po_allocations a JOIN po_lines p ON p.po_line_id=a.po_line_id
                           JOIN indent_lines i ON i.indent_line_id=a.indent_line_id
                           JOIN indent_headers h ON h.indent_header_id=i.indent_header_id"""
                allocations = database.fetch_dataframe(query)
                self.assertFalse(allocations.empty)
                self.assertNotIn("345A", allocations["po_no"].tolist())
                for _, row in allocations.iterrows():
                    self.assertTrue(_po_facility_matches(row.institution_name, row.indent_district, row.location, row.district))
                for _, group in allocations.groupby("po_line_id"):
                    self.assertLessEqual(group.allocated_qty.astype(float).sum(), float(group.iloc[0].ordered_qty))
                expected = {
                    "Infusion Pump": "215A", "Syrange Pump": "215A", "Crashcart": "211A",
                    "Auto Clave HP Vertical (2 bin)": "221A", "Operation Table-Hydraulic Major": "217A",
                    "Oxygen Cylinder Trolley(SS)": "217A",
                    "100 M.A. X-ray machine Mobile HF": "209A", "200 M.A. X-ray machine HF": "209A",
                    "Dental X ray machine": "210A", "Monitors ( 5 Para)": "214A", "ECG machine": "214A",
                    "CR Machine": "216A", "Wheel Chair (SS)": "217A",
                    "Hospital Cots Pediatric (ISI Model) with mattresses": "217A",
                    "Stools Revolving ;": "218A", "Office Table": "218A", "Stool": "218A",
                    "Racks -Steel - Wooden": "218A", "Surgical Diathermy": "219A",
                    "Shadowless Lamp stand model": "219A", "Bed side screen (SS- Godrej Model)": "220A",
                    "Stretcher / Patient Trolley (SS)": "223A", "Office Chairs": "224A",
                    "Instrument Trolley (SS)": "224A", "Portable Scanning Machine": "213A",
                }
                for item, po in expected.items():
                    self.assertIn(po, allocations.loc[allocations.source_item_name == item, "po_no"].tolist())
                expected_multi = {
                    "Hospital Cots (ISI Model) with mattresses": {"217A", "224A"},
                    "Saline Stand (SS)": {"220A", "223A"},
                }
                for item, po_set in expected_multi.items():
                    self.assertEqual(
                        set(allocations.loc[allocations.source_item_name == item, "po_no"].tolist()), po_set
                    )
                # CR machine's accessory plates are linked as package components, not
                # reported as orphan PO lines, and do not inflate its confirmed quantity.
                cr_lifecycle = database.fetch_dataframe(
                    "SELECT matched_po_ids, confirmed_ordered_qty FROM indent_lifecycle "
                    "WHERE source_item_name='CR Machine'"
                )
                self.assertEqual(len(json.loads(cr_lifecycle.iloc[0]["matched_po_ids"])), 3)
                self.assertEqual(cr_lifecycle.iloc[0]["confirmed_ordered_qty"], "1")
                # Probable matches (needing manual confirmation) and genuine PO-side
                # orphans/data-quality issues are recorded as exceptions, not silently dropped.
                exceptions = database.fetch_dataframe("SELECT exception_code, summary FROM exceptions")
                orphan_summaries = " ".join(
                    exceptions.loc[exceptions.exception_code == "PO_LINE_NOT_FOUND_IN_INDENT", "summary"]
                )
                for expected_orphan in ["3 Seater Visitor Bench"]:
                    self.assertIn(expected_orphan, orphan_summaries)
                # "Foot step" is a human-curated probable match (not a generic fuzzy
                # hit), so it must surface as a review candidate instead of an orphan.
                self.assertNotIn("Foot step", orphan_summaries)
                foot_step_payload = database.fetch_dataframe(
                    "SELECT raw_payload_json FROM indent_lifecycle WHERE source_item_name='Double & Single Step Stools'"
                )
                self.assertIn(
                    "Foot step",
                    json.dumps(json.loads(foot_step_payload.iloc[0]["raw_payload_json"]).get("po_review_candidates", [])),
                )
                amount_mismatches = exceptions.loc[exceptions.exception_code == "PO_HEADER_AMOUNT_MISMATCH", "summary"]
                self.assertTrue(amount_mismatches.str.contains("223A").any())
                before = len(database.fetch_dataframe("SELECT * FROM indent_lifecycle"))
                second = run_inventory_scan(config, database)
                self.assertEqual(second.files_ingested, 0)
                self.assertEqual(len(database.fetch_dataframe("SELECT * FROM indent_lifecycle")), before)
                self.assertEqual(len(database.fetch_dataframe(query)), len(allocations))
                output = build_report(database, config)
                workbook = load_workbook(output)
                template = load_workbook(ROOT / "Indent_Tracker.xlsx")
                ws = workbook["Indent Tracker"]
                # The approved template sheets must stay intact and in place; the
                # reconciliation adds an extra "beautiful report" sheet alongside them.
                self.assertEqual(workbook.sheetnames[: len(template.sheetnames)], template.sheetnames)
                self.assertIn("Summary Report", workbook.sheetnames)
                self.assertNotIn("Exceptions", workbook.sheetnames)
                self.assertEqual(ws.freeze_panes, "A4")
                self.assertEqual(str(ws.merged_cells), str(template["Indent Tracker"].merged_cells))
                self.assertEqual(len(ws.data_validations.dataValidation), 7)
                output_rows = list(ws.iter_rows(min_row=4, values_only=True))
                for item, po in expected.items():
                    self.assertTrue(any(row[7] == item and po in (row[16] or "") for row in output_rows))
                for item, po_set in expected_multi.items():
                    self.assertTrue(any(row[7] == item and po_set.issubset(set((row[16] or "").split(", "))) for row in output_rows))
                self.assertFalse(any("345A" in (row[16] or "") for row in output_rows))
                self.assertTrue(any(cell.comment and "Probable Match" in cell.comment.text for cell in ws["Q"]))
                workbook.close()
                template.close()
            finally:
                database.close()


if __name__ == "__main__":
    unittest.main()