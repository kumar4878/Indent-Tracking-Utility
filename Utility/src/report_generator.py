from __future__ import annotations

import json
import shutil
import warnings
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database_service import ReconciliationDatabase


TEMPLATE_FILE_NAME = "Indent_Tracker.xlsx"
TRACKER_SHEET_NAME = "Indent Tracker"
TRACKER_HEADERS = [
    "Sl. No",
    "Inward No",
    "Indent No",
    "District",
    "Hospital",
    "HOD",
    "Indent Date",
    "Equipment Name",
    "Equipment Type",
    "Quantity",
    "Prioritization",
    "Budget Sanction",
    "Case Worker",
    "Rate Contract",
    "Tender No",
    "Tender Stage",
    "PO No",
    "Supplied Qty",
    "Pending Qty",
    "Supplied",
]
PRIORITIZATION_OPTIONS = {
    "Required",
    "Not Required",
    "Reprioritization Sent - Awaiting Reply",
}
# Diagnostic-only column appended after the approved template's 20 columns.
# It never feeds dropdowns/validations and is recomputed fresh on every export.
MATCH_STATUS_HEADER = "Match Status / Closest Candidate (for review only)"
MATCH_STATUS_COLUMN = len(TRACKER_HEADERS) + 1
CONFIDENCE_SCORE_HEADER = "Confidence Score 0-100 (diagnostic only)"
CONFIDENCE_SCORE_COLUMN = MATCH_STATUS_COLUMN + 1

SUMMARY_SHEET_NAME = "Summary Report"

REPORT_TITLE_FONT = Font(bold=True, size=16, color="1F4E78")
REPORT_SUBTITLE_FONT = Font(italic=True, size=10, color="6B7280")
SECTION_FONT = Font(bold=True, size=12, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="1F4E78")
TABLE_HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
TABLE_HEADER_FILL = PatternFill("solid", fgColor="2E75B6")
BAND_FILL = PatternFill("solid", fgColor="EAF1FB")
METRIC_VALUE_FONT = Font(bold=True, size=13, color="1F4E78")
_THIN_SIDE = Side(style="thin", color="B7C6D9")
TABLE_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)


def _copy_cell_style(source, target) -> None:
    """Mirror an existing template cell's font/border/alignment/fill so new
    content blends in instead of falling back to openpyxl's plain defaults.
    """
    target.font = copy(source.font)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.fill = copy(source.fill)
    target.number_format = source.number_format


def _clean_value(value: Any) -> Any:
    if value is None:
        return ""
    if pd.isna(value):
        return ""
    if isinstance(value, str) and value.strip().lower() in {"nan", "none", "null"}:
        return ""
    return value


def _parse_id_list(value: Any) -> List[int]:
    text = str(_clean_value(value)).strip()
    if not text:
        return []
    try:
        raw_values = json.loads(text)
    except json.JSONDecodeError:
        return []
    ids: List[int] = []
    for raw_value in raw_values:
        try:
            ids.append(int(raw_value))
        except (TypeError, ValueError):
            continue
    return ids


def _frame_lookup(frame: pd.DataFrame, id_column: str) -> Dict[int, Dict[str, Any]]:
    if frame.empty:
        return {}
    lookup: Dict[int, Dict[str, Any]] = {}
    for row in frame.to_dict(orient="records"):
        try:
            lookup[int(row[id_column])] = row
        except (KeyError, TypeError, ValueError):
            continue
    return lookup


def _first_lookup_value(ids: Iterable[int], lookup: Dict[int, Dict[str, Any]]) -> Dict[str, Any]:
    for record_id in ids:
        if record_id in lookup:
            return lookup[record_id]
    return {}


def _join_po_numbers(ids: Iterable[int], po_lookup: Dict[int, Dict[str, Any]]) -> str:
    values: List[str] = []
    seen = set()
    for record_id in ids:
        po_record = po_lookup.get(record_id, {})
        po_no = str(_clean_value(po_record.get("po_no", ""))).strip()
        if po_no and po_no not in seen:
            seen.add(po_no)
            values.append(po_no)
    return ", ".join(values)


def _join_review_po_numbers(reviews: List[Dict[str, Any]]) -> str:
    """PO numbers for not-yet-confirmed candidates (score 70-90%+ or an
    analyst-curated probable pair). Shown in the PO No. column so reviewers
    see the candidate directly, while the remarks column marks it Probable.
    """
    values: List[str] = []
    seen = set()
    for review in reviews:
        po_no = str(_clean_value(review.get("po_no", ""))).strip()
        if po_no and po_no not in seen:
            seen.add(po_no)
            values.append(po_no)
    return ", ".join(values)


def _derive_rc_status(rc_record: Dict[str, Any]) -> str:
    if not rc_record:
        return "Not Available"
    valid_upto = str(_clean_value(rc_record.get("rc_valid_upto", ""))).strip()
    if not valid_upto:
        return "Validity Under Verification"
    use_dayfirst = "-" not in valid_upto[:10]
    parsed = pd.to_datetime(valid_upto, errors="coerce", dayfirst=use_dayfirst)
    if pd.isna(parsed):
        return "Validity Under Verification"
    return "Expired" if parsed.date() < datetime.now().date() else "Available"


def _derive_prioritization(lifecycle_status: Any) -> str:
    status = str(_clean_value(lifecycle_status)).strip()
    if status in PRIORITIZATION_OPTIONS:
        return status
    if status == "Returned for Reprioritization":
        return "Reprioritization Sent - Awaiting Reply"
    return "Required"


def _clear_tracker_rows(worksheet) -> None:
    for row in worksheet.iter_rows(min_row=4, max_row=worksheet.max_row, min_col=1, max_col=CONFIDENCE_SCORE_COLUMN):
        for cell in row:
            cell.value = None
            cell.comment = None


def _match_status_text(
    lifecycle_row: Dict[str, Any],
    payload: Dict[str, Any],
    rc_ids: List[int],
    tender_ids: List[int],
    po_ids: List[int],
) -> str:
    """Summarize, in plain language, how this indent line's PO No. cell was
    decided: a fully/partially confirmed PO link, the single closest
    not-yet-confirmed PO candidate (with its similarity score) when one
    exists, or an explicit 'not matched' reason (RC available, tender in
    progress, or nothing found at all).
    """
    status = str(_clean_value(lifecycle_row.get("current_lifecycle_status", "")))
    reviews = payload.get("po_review_candidates", []) or []
    suggested_components = payload.get("suggested_components", []) or []
    suggestion_suffix = (
        f"; {len(suggested_components)} suggested component(s) - Confirmation Required" if suggested_components else ""
    )

    if po_ids and status in {"Fully Ordered", "Partially Ordered"}:
        label = "Confirmed Match" if status == "Fully Ordered" else "Confirmed Match (Partial Quantity)"
        if reviews:
            label += f"; {len(reviews)} additional unconfirmed candidate(s) also exist"
        return label + suggestion_suffix

    if reviews:
        top = reviews[0]
        score = top.get("match_score")
        if isinstance(score, (int, float)) and not pd.isna(score):
            score_text = f", similarity {int(score)}%"
        else:
            score_text = ", analyst-flagged candidate"
        qty_text = f" qty {top.get('ordered_qty')}" if top.get("ordered_qty") not in (None, "") else ""
        return (
            f"Probable Match - Closest: PO {top.get('po_no', '')} "
            f"'{top.get('source_item_name', '')}'{qty_text}{score_text} - Confirmation Required"
        )

    if rc_ids:
        return "Not Matched to PO - Rate Contract Available (PO Pending)"
    if tender_ids:
        return "Not Matched to PO - Tendering In Progress"
    return "Not Matched - No RC/Tender/PO Candidate Found"


def _annotate_po_cells(database: ReconciliationDatabase, worksheet) -> None:
    """Expose allocation evidence/review reasons without changing template columns."""
    lifecycle = database.fetch_dataframe(
        "SELECT matched_po_ids, raw_payload_json, confirmed_ordered_qty, po_pending_qty FROM indent_lifecycle ORDER BY lifecycle_id"
    )
    po_lookup = _frame_lookup(database.fetch_dataframe("SELECT * FROM po_lines"), "po_line_id")
    allocations = database.fetch_dataframe("SELECT po_line_id, allocation_status FROM po_allocations")
    status_by_po_id: Dict[int, str] = {}
    for row in allocations.to_dict(orient="records"):
        try:
            status_by_po_id[int(row["po_line_id"])] = str(row.get("allocation_status") or "")
        except (TypeError, ValueError):
            continue

    for row_index, record in enumerate(lifecycle.to_dict(orient="records"), start=4):
        matched = _parse_id_list(record.get("matched_po_ids"))
        notes = []
        for po_id in matched:
            po = po_lookup.get(po_id)
            if not po:
                continue
            label = "Package Component" if status_by_po_id.get(po_id) == "Package Component" else "Confirmed"
            notes.append(
                f"[{label}] PO {po['po_no']}: {po['source_item_name']}; {po['location']}, {po['district']}. "
                f"Source: {po['source_file']} / {po['source_sheet']} / row {po['source_row']}."
            )
        if matched:
            notes.append(f"Allocated ordered quantity: {record['confirmed_ordered_qty']}. This is not evidence of supply/delivery.")
        else:
            notes.append("No confirmed PO allocation. A matching hospital/district and a sufficiently strong item match with available quantity are required.")
        payload = json.loads(record.get("raw_payload_json") or "{}")
        reviews = payload.get("po_review_candidates", [])
        if reviews:
            notes.append("Match Status: Probable Match - Action: Technical/Administrative confirmation required.")
            notes.extend(
                f"PO {po.get('po_no', '')}: {po.get('source_item_name', '')} (qty {po.get('ordered_qty', '')})"
                for po in reviews
            )
        worksheet.cell(row_index, 17).comment = Comment("\n".join(notes), "Reconciliation Utility")


def _ensure_template_headers(worksheet) -> None:
    headers = [worksheet.cell(row=3, column=index).value for index in range(1, len(TRACKER_HEADERS) + 1)]
    if headers != TRACKER_HEADERS:
        raise ValueError("Indent_Tracker.xlsx does not match the expected tracker template headers.")


def _build_tracker_rows(database: ReconciliationDatabase) -> List[List[Any]]:
    lifecycle = database.fetch_dataframe(
        "SELECT lifecycle_id, indent_line_id, indent_header_id, source_item_name, required_qty, matched_rc_ids, matched_tender_ids, matched_po_ids, confirmed_ordered_qty, po_pending_qty, current_lifecycle_status, confidence_score, raw_payload_json FROM indent_lifecycle ORDER BY lifecycle_id"
    )
    if lifecycle.empty:
        return []

    headers = database.fetch_dataframe(
        "SELECT indent_header_id, indent_no, indent_date, institution_name, district, as_amount, department FROM indent_headers"
    )
    tenders = database.fetch_dataframe(
        "SELECT tender_item_id, tender_no, tender_stage, case_worker FROM tender_items"
    )
    po_lines = database.fetch_dataframe(
        "SELECT po_line_id, po_no, case_worker_name FROM po_lines"
    )
    rc_items = database.fetch_dataframe(
        "SELECT rc_item_id, rc_valid_upto FROM rc_items"
    )

    header_lookup = _frame_lookup(headers, "indent_header_id")
    tender_lookup = _frame_lookup(tenders, "tender_item_id")
    po_lookup = _frame_lookup(po_lines, "po_line_id")
    rc_lookup = _frame_lookup(rc_items, "rc_item_id")

    rows: List[List[Any]] = []
    for row_number, lifecycle_row in enumerate(lifecycle.to_dict(orient="records"), start=1):
        header = header_lookup.get(int(lifecycle_row["indent_header_id"]), {})
        rc_ids = _parse_id_list(lifecycle_row.get("matched_rc_ids"))
        tender_ids = _parse_id_list(lifecycle_row.get("matched_tender_ids"))
        po_ids = _parse_id_list(lifecycle_row.get("matched_po_ids"))

        rc_record = _first_lookup_value(rc_ids, rc_lookup)
        tender_record = _first_lookup_value(tender_ids, tender_lookup)
        po_record = _first_lookup_value(po_ids, po_lookup)
        case_worker = str(_clean_value(tender_record.get("case_worker", ""))).strip()
        if not case_worker:
            case_worker = str(_clean_value(po_record.get("case_worker_name", ""))).strip()
        payload = json.loads(lifecycle_row.get("raw_payload_json") or "{}")
        match_status_text = _match_status_text(lifecycle_row, payload, rc_ids, tender_ids, po_ids)
        po_no_display = _join_po_numbers(po_ids, po_lookup)
        if not po_no_display:
            po_no_display = _join_review_po_numbers(payload.get("po_review_candidates", []) or [])

        rows.append([
            row_number,
            "",
            _clean_value(header.get("indent_no", "")),
            _clean_value(header.get("district", "")),
            _clean_value(header.get("institution_name", "")),
            _clean_value(header.get("department", "")),
            _clean_value(header.get("indent_date", "")),
            _clean_value(lifecycle_row.get("source_item_name", "")),
            "",
            _clean_value(lifecycle_row.get("required_qty", "")),
            _derive_prioritization(lifecycle_row.get("current_lifecycle_status", "")),
            _clean_value(header.get("as_amount", "")),
            case_worker,
            _derive_rc_status(rc_record),
            _clean_value(tender_record.get("tender_no", "")),
            _clean_value(tender_record.get("tender_stage", "")),
            po_no_display,
            "",
            _clean_value(lifecycle_row.get("po_pending_qty", "")),
            "",
            match_status_text,
            _clean_value(lifecycle_row.get("confidence_score", "")),
        ])

    return rows


def _summary_counts(database: ReconciliationDatabase) -> Dict[str, int]:
    lifecycle = database.fetch_dataframe(
        "SELECT matched_rc_ids, matched_tender_ids, matched_po_ids, current_lifecycle_status, raw_payload_json FROM indent_lifecycle"
    )
    counts = {
        "Total Indent Lines": int(len(lifecycle)),
        "Confirmed Match": 0,
        "Probable Match (Needs Review)": 0,
        "Rate Contract Available (PO Pending)": 0,
        "Tendering In Progress": 0,
        "Not Matched - No Candidate Found": 0,
    }
    for record in lifecycle.to_dict(orient="records"):
        rc_ids = _parse_id_list(record.get("matched_rc_ids"))
        tender_ids = _parse_id_list(record.get("matched_tender_ids"))
        po_ids = _parse_id_list(record.get("matched_po_ids"))
        payload = json.loads(record.get("raw_payload_json") or "{}")
        text = _match_status_text(record, payload, rc_ids, tender_ids, po_ids)
        if text.startswith("Confirmed Match"):
            counts["Confirmed Match"] += 1
        elif text.startswith("Probable Match"):
            counts["Probable Match (Needs Review)"] += 1
        elif "Rate Contract Available" in text:
            counts["Rate Contract Available (PO Pending)"] += 1
        elif "Tendering In Progress" in text:
            counts["Tendering In Progress"] += 1
        else:
            counts["Not Matched - No Candidate Found"] += 1
    return counts


def _hospital_breakdown(database: ReconciliationDatabase) -> List[List[Any]]:
    lifecycle = database.fetch_dataframe(
        "SELECT indent_header_id, matched_rc_ids, matched_tender_ids, matched_po_ids, current_lifecycle_status, raw_payload_json FROM indent_lifecycle"
    )
    headers = database.fetch_dataframe("SELECT indent_header_id, institution_name, district FROM indent_headers")
    header_lookup = _frame_lookup(headers, "indent_header_id")

    buckets: Dict[int, Dict[str, int]] = {}
    for record in lifecycle.to_dict(orient="records"):
        header_id = int(record.get("indent_header_id") or 0)
        bucket = buckets.setdefault(header_id, {"total": 0, "confirmed": 0, "probable": 0, "not_matched": 0})
        bucket["total"] += 1
        rc_ids = _parse_id_list(record.get("matched_rc_ids"))
        tender_ids = _parse_id_list(record.get("matched_tender_ids"))
        po_ids = _parse_id_list(record.get("matched_po_ids"))
        payload = json.loads(record.get("raw_payload_json") or "{}")
        text = _match_status_text(record, payload, rc_ids, tender_ids, po_ids)
        if text.startswith("Confirmed Match"):
            bucket["confirmed"] += 1
        elif text.startswith("Probable Match"):
            bucket["probable"] += 1
        else:
            bucket["not_matched"] += 1

    rows: List[List[Any]] = []
    for header_id, bucket in buckets.items():
        header = header_lookup.get(header_id, {})
        rows.append([
            _clean_value(header.get("district", "")),
            _clean_value(header.get("institution_name", "")),
            bucket["total"],
            bucket["confirmed"],
            bucket["probable"],
            bucket["not_matched"],
        ])
    rows.sort(key=lambda row: (str(row[0]), str(row[1])))
    return rows


def _style_title(worksheet, row: int, span: int, text: str) -> None:
    cell = worksheet.cell(row=row, column=1, value=text)
    cell.font = REPORT_TITLE_FONT
    worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)


def _style_subtitle(worksheet, row: int, span: int, text: str) -> None:
    cell = worksheet.cell(row=row, column=1, value=text)
    cell.font = REPORT_SUBTITLE_FONT
    worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)


def _style_section_header(worksheet, row: int, span: int, text: str) -> None:
    worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = worksheet.cell(row=row, column=1, value=text)
    for column_index in range(1, span + 1):
        target = worksheet.cell(row=row, column=column_index)
        target.fill = SECTION_FILL
        target.font = SECTION_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[row].height = 20


def _write_table(worksheet, start_row: int, headers: List[str], rows: List[List[Any]], highlight_column: int = -1) -> int:
    for column_index, header_text in enumerate(headers, start=1):
        cell = worksheet.cell(row=start_row, column=column_index, value=header_text)
        cell.font = TABLE_HEADER_FONT
        cell.fill = TABLE_HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = TABLE_BORDER

    current_row = start_row + 1
    for row_offset, row_values in enumerate(rows):
        for column_index, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row=current_row, column=column_index, value=value)
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(
                horizontal="center" if column_index != 1 else "left",
                vertical="center",
                wrap_text=column_index == len(headers),
            )
            if row_offset % 2 == 1:
                cell.fill = BAND_FILL
            if column_index - 1 == highlight_column:
                cell.font = METRIC_VALUE_FONT
        current_row += 1
    return current_row


def _build_summary_sheet(workbook, database: ReconciliationDatabase) -> None:
    if SUMMARY_SHEET_NAME in workbook.sheetnames:
        del workbook[SUMMARY_SHEET_NAME]
    worksheet = workbook.create_sheet(SUMMARY_SHEET_NAME)

    _style_title(worksheet, 1, 6, "TGMSIDC Indent-to-Procurement Reconciliation - Summary Report")
    _style_subtitle(worksheet, 2, 6, f"Generated on {datetime.now().strftime('%d-%b-%Y %H:%M')}")

    counts = _summary_counts(database)
    exceptions = database.fetch_dataframe("SELECT exception_code FROM exceptions")
    orphan_count = int((exceptions["exception_code"] == "PO_LINE_NOT_FOUND_IN_INDENT").sum()) if not exceptions.empty else 0
    mismatch_count = int((exceptions["exception_code"] == "PO_HEADER_AMOUNT_MISMATCH").sum()) if not exceptions.empty else 0

    _style_section_header(worksheet, 4, 2, "Key Metrics")
    metric_rows = [[label, value] for label, value in counts.items()]
    metric_rows.append(["Orphan PO Lines (No Indent Match)", orphan_count])
    metric_rows.append(["PO Header/Line Amount Mismatches", mismatch_count])
    next_row = _write_table(worksheet, 5, ["Metric", "Count"], metric_rows, highlight_column=1)

    section_row = next_row + 1
    _style_section_header(worksheet, section_row, 6, "Breakdown by District / Hospital")
    breakdown_rows = _hospital_breakdown(database)
    _write_table(
        worksheet,
        section_row + 1,
        ["District", "Hospital", "Total Lines", "Confirmed", "Probable", "Not Matched"],
        breakdown_rows,
    )

    widths = {1: 34, 2: 30, 3: 14, 4: 14, 5: 14, 6: 16}
    for column_index, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    worksheet.sheet_view.showGridLines = False


def build_report(database: ReconciliationDatabase, config: Dict[str, Any]) -> Path:
    root_path = Path(str(config["root_path"]))
    output_folder = root_path / str(config["output_folder"])
    output_folder.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_folder / f"TGMSIDC_Indent_Lifecycle_{timestamp}.xlsx"
    latest_path = output_folder / "TGMSIDC_Indent_Lifecycle_Latest.xlsx"
    template_path = root_path.parent / TEMPLATE_FILE_NAME

    if not template_path.exists():
        raise FileNotFoundError(f"Tracker template not found: {template_path}")

    workbook = load_workbook(template_path)
    worksheet = workbook[TRACKER_SHEET_NAME]
    _ensure_template_headers(worksheet)
    _clear_tracker_rows(worksheet)

    header_style_source = worksheet.cell(row=3, column=len(TRACKER_HEADERS))
    match_status_header_cell = worksheet.cell(row=3, column=MATCH_STATUS_COLUMN, value=MATCH_STATUS_HEADER)
    _copy_cell_style(header_style_source, match_status_header_cell)
    confidence_header_cell = worksheet.cell(row=3, column=CONFIDENCE_SCORE_COLUMN, value=CONFIDENCE_SCORE_HEADER)
    _copy_cell_style(header_style_source, confidence_header_cell)

    tracker_rows = _build_tracker_rows(database)
    for row_index, row_values in enumerate(tracker_rows, start=4):
        for column_index, value in enumerate(row_values, start=1):
            worksheet.cell(row=row_index, column=column_index, value=value)
        # Mirror whatever font/border this row already uses (template rows and
        # any rows beyond the template's pre-styled range) so the new column
        # blends in instead of using openpyxl's plain default style.
        row_style_source = worksheet.cell(row=row_index, column=len(TRACKER_HEADERS))
        _copy_cell_style(row_style_source, worksheet.cell(row=row_index, column=MATCH_STATUS_COLUMN))
        _copy_cell_style(row_style_source, worksheet.cell(row=row_index, column=CONFIDENCE_SCORE_COLUMN))
    _annotate_po_cells(database, worksheet)

    last_row = max(worksheet.max_row, len(tracker_rows) + 3)
    last_column_letter = worksheet.cell(row=1, column=CONFIDENCE_SCORE_COLUMN).column_letter
    worksheet.auto_filter.ref = f"A1:{last_column_letter}{last_row}"

    _build_summary_sheet(workbook, database)

    workbook.save(output_path)

    if bool(config.get("create_latest_output", True)):
        try:
            shutil.copy2(output_path, latest_path)
        except PermissionError:
            warnings.warn(
                f"Report saved to {output_path}, but {latest_path} is open or not writable. "
                "The Latest workbook was NOT updated; open the timestamped report instead.",
                RuntimeWarning,
                stacklevel=2,
            )

    return output_path
